"""K-apt 면적정보 xlsx 로더.

필요 패키지: openpyxl
"""

import pathlib
import sys

import openpyxl


HERE = pathlib.Path(__file__).resolve().parent
ROOT = HERE.parent
KAPT_AREA_XLSX_PATH = ROOT / "data" / "kapt-area-info.xlsx"
COMPLEX_AREA_PATH = ROOT / "data" / "complex-area.csv"
MATCHING_FAILURES_PATH = ROOT / "data" / "matching-failures.csv"

HEADER_NAME = "단지코드"
REQUIRED_COLUMNS = (HEADER_NAME, "주거전용면적(세부)", "세대수")
TEXT_VALUE_COLUMNS = REQUIRED_COLUMNS


def load_kapt_area_rows(path=KAPT_AREA_XLSX_PATH) -> list[dict]:
    """헤더 행을 찾아 다음 행부터 원본 값의 dict 목록으로 반환."""
    workbook = openpyxl.load_workbook(path, read_only=True)
    rows = None
    try:
        if len(workbook.sheetnames) != 1:
            raise ValueError(
                "K-apt 면적정보 xlsx의 시트가 1개여야 합니다: "
                f"{workbook.sheetnames!r}"
            )

        rows = workbook.active.iter_rows(values_only=True)
        for header_row_number, values in enumerate(rows, start=1):
            if HEADER_NAME in values:
                headers = values
                break
        else:
            raise ValueError(f"필수 헤더 {HEADER_NAME!r}를 찾지 못했습니다.")

        missing = [name for name in REQUIRED_COLUMNS if name not in headers]
        if missing:
            raise ValueError(f"필수 열이 없습니다: {', '.join(missing)}")

        data = []
        for row_number, values in enumerate(rows, start=header_row_number + 1):
            record = dict(zip(headers, values))
            for name in TEXT_VALUE_COLUMNS:
                value = record[name]
                if value is not None and not isinstance(value, str):
                    raise ValueError(
                        f"{name!r} 값은 문자열이어야 합니다: {row_number}행 {value!r}"
                    )
            data.append(record)

        if not data:
            raise ValueError("K-apt 면적정보에 데이터 행이 없습니다.")
        return data
    finally:
        if rows is not None:
            rows.close()
        workbook.close()


def main(
    path=KAPT_AREA_XLSX_PATH,
    *,
    out_area=COMPLEX_AREA_PATH,
    out_failures=MATCHING_FAILURES_PATH,
) -> None:
    """면적정보를 검증해 읽고 기본 건수를 출력.

    출력 경로는 다음 스텝에서 사용한다.
    지금부터 인자로 받아 CSV 쓰기 테스트를 격리한다.
    """
    _ = out_area, out_failures
    try:
        rows = load_kapt_area_rows(path)
    except ValueError as exc:
        sys.exit(f"[중단] {exc}")

    blank_codes = sum(row[HEADER_NAME] is None for row in rows)
    distinct_codes = len({row[HEADER_NAME] for row in rows if row[HEADER_NAME] is not None})
    print(f"행 수: {len(rows):,}")
    print(f"{HEADER_NAME} 빈 행: {blank_codes:,}")
    print(f"{HEADER_NAME} distinct: {distinct_codes:,}")


if __name__ == "__main__":
    main()
